"""Helpers de formato para los reportes Excel de la Fase 1."""

from openpyxl.styles import Font, PatternFill, Alignment, Border, Side


def _thin():
    s = Side(style="thin", color="BDBDBD")
    return Border(left=s, right=s, top=s, bottom=s)


def aplicar_header(ws, color_hex, text_color="FFFFFF"):
    for cell in ws[1]:
        cell.font      = Font(name="Arial", bold=True, size=10, color=text_color)
        cell.fill      = PatternFill("solid", fgColor=color_hex)
        cell.alignment = Alignment(horizontal="center", vertical="center",
                                   wrap_text=True)
    ws.row_dimensions[1].height = 28


def aplicar_fila(ws, color_hex, wrap=True):
    fila = ws.max_row
    for c in range(1, ws.max_column + 1):
        cell = ws.cell(fila, c)
        cell.fill      = PatternFill("solid", fgColor=color_hex)
        cell.font      = Font(name="Arial", size=10)
        cell.alignment = Alignment(vertical="center", wrap_text=wrap)
        cell.border    = _thin()


def ancho(ws, letra, valor):
    ws.column_dimensions[letra].width = valor
