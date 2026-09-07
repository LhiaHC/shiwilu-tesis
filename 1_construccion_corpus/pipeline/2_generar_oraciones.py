"""
Etapa 2 — Generacion de oraciones para las categorias con deficit (OE1).

Genera con la API las oraciones en espanol que faltan para llegar a 100 pares
por intencion y las deja listas para que el hablante nativo las traduzca.

Entrada: 1_construccion_corpus/intermedios/1_corpus_pares_anotados.xlsx
Salida:  1_construccion_corpus/intermedios/2_oraciones_para_traducir.xlsx
         1_construccion_corpus/intermedios/logs/2_oraciones_generadas_log.json

Requiere ANTHROPIC_API_KEY. Una vez traducido, el archivo se guarda como
2_oraciones_traducidas.xlsx, que es lo que consume la Etapa 3.
"""

import os
import json
import re
import sys
from datetime import datetime

from dotenv import load_dotenv
load_dotenv()

import anthropic
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import PatternFill

import _ruta_raiz  # noqa: F401  (deja importable el paquete `shiwilu`)
from shiwilu.dominios import DOMINIOS
from shiwilu.excel import ancho, aplicar_fila, aplicar_header
from shiwilu.rutas import (
    INTERMEDIOS, LOGS, PARES_ANOTADOS, preparar_directorios,
)
from shiwilu.taxonomia import COLOR_INT, DESC_INT

preparar_directorios()

API_KEY = os.environ.get("ANTHROPIC_API_KEY", "")
if not API_KEY:
    raise ValueError("[ERROR] No se encontro ANTHROPIC_API_KEY en el archivo .env")

MODELO   = "claude-sonnet-4-6"
OBJETIVO = 100
LONG_MIN = 1
LONG_MAX = 5

EJEMPLOS_REALES = {
    "SAL": ["Hola.", "Hasta luego.", "Nos vemos.", "Bienvenido.", "Buenas tardes.",
            "Buenos dias.", "Adios.", "Chau.", "Bienvenidos.", "Buenas noches."],
    "EMO": ["Que divertido!", "Gracias.", "Que bonito!", "Que lastima!", "Que pena!",
            "Lo siento.", "Que susto!", "Que humillante!", "Que interesante!", "Que aburrimiento!"],
    "AFI": ["Esta perfecto.", "Es correcto.", "Esta bien.", "De acuerdo.",
            "Esta listo.", "Claro.", "Exacto.", "Asi es.", "Ya.", "Por supuesto."],
    "NEG": ["Ni idea.", "No entiendo.", "Me niego.", "Nadie llamo.", "No se.",
            "Nunca hablamos.", "Nadie escucha.", "Imposible.", "Nunca.", "Sin agua."],
}

RESTRICCIONES_PATRON = {
    "EMO": (
        "Distribuye las oraciones usando TODOS estos patrones (ningun patron puede superar el 25%):\n"
        "- Exclamacion directa: Bravo! / Lo siento! / Que lastima! / Que horror!\n"
        "- Ay + emocion: Ay, que tristeza! / Ay, no!\n"
        "- Verbo emocional: Me alegra. / Me duele. / Me da pena. / Me da gusto. / Me entristece.\n"
        "- Cuanto/Como: Cuanto me alegra! / Como duele! / Cuanto lo extrano!\n"
        "- Que + adjetivo/sustantivo (MAX 25% del total): Que bonito! / Que pena!\n"
        "PROHIBIDO: mas del 25% de oraciones que empiecen con 'Que' o '!Que'."
    ),
    "AFI": (
        "Distribuye las oraciones usando TODOS estos patrones (ningun patron puede superar el 25%):\n"
        "- Palabra sola: Claro. / Exacto. / Correcto. / Verdad. / Perfecto. / Ya. / Listo.\n"
        "- Frase corta: De acuerdo. / Asi es. / Esta bien. / Por supuesto. / Asi mismo. / Entendido.\n"
        "- Confirmacion directa: Eso es. / Eso mismo. / Exactamente. / Acepto. / Confirmado.\n"
        "- Si + frase (MAX 25% del total): Si. / Si, claro. / Si, ya. / Si, correcto.\n"
        "PROHIBIDO: mas del 25% de oraciones que empiecen con 'Si' o 'Si,'."
    ),
    "NEG": (
        "Distribuye las oraciones usando TODOS estos patrones (ningun patron puede superar el 30%):\n"
        "- Palabra sola: Nunca. / Jamas. / Imposible. / Nada. / Nadie.\n"
        "- Sin + sustantivo: Sin agua. / Sin fuego. / Sin yuca. / Sin peces. / Sin camino.\n"
        "- Me + verbo: Me niego. / Me falta agua. / Me cuesta.\n"
        "- Nadie + verbo: Nadie llega. / Nadie sabe. / Nadie viene. / Nadie escucha.\n"
        "- Nunca/Jamas + verbo: Nunca pesco. / Jamas voy. / Nunca llego.\n"
        "- Ya no / Se acabo: Ya no hay. / Ya se fue. / Se acabo. / Se perdio.\n"
        "- No + verbo (MAX 30%): No se. / No puedo. / No quiero. / No hay.\n"
        "PROHIBIDO: mas del 30% de oraciones que empiecen con la palabra 'No'."
    ),
}


def _norm(s: str) -> str:
    return s.strip().lower().strip(".,!?¡¿;: ")


def leer_corpus_existente() -> set:
    if not PARES_ANOTADOS.exists():
        return set()
    df      = pd.read_excel(PARES_ANOTADOS, sheet_name=0)
    col_esp = next(c for c in df.columns if "spa" in c.lower() or "spanol" in c.lower())
    return set(df[col_esp].dropna().apply(lambda x: _norm(str(x))))


def leer_deficit() -> dict:
    if not PARES_ANOTADOS.exists():
        raise FileNotFoundError(
            f"No se encontro {PARES_ANOTADOS.name}. Ejecuta primero: "
            "python 1_construccion_corpus/pipeline/1_generar_pares_anotados.py"
        )
    xl   = pd.ExcelFile(PARES_ANOTADOS)
    hoja = next((s for s in xl.sheet_names if "istrib" in s), xl.sheet_names[1])
    df   = pd.read_excel(PARES_ANOTADOS, sheet_name=hoja)

    col_int = next(c for c in df.columns if "nten" in c.lower())
    col_def = next(c for c in df.columns if "ficit" in c.lower() or "eficit" in c.lower())

    deficit = {}
    for _, row in df.iterrows():
        intent = str(row[col_int]).strip()
        val    = row[col_def]
        if str(val) not in ("—", "nan", "") and float(val) > 0:
            deficit[intent] = int(float(val))
    return deficit


def vocab_para_prompt() -> str:
    lineas = []
    for dom, info in DOMINIOS.items():
        semillas = ", ".join(info["semillas"][:15])
        lineas.append(f"  {dom} ({info['desc']}): {semillas}...")
    return "\n".join(lineas)


def generar_oraciones(categoria: str, cantidad: int,
                      vocab_str: str, excluir: set) -> list:
    client = anthropic.Anthropic(api_key=API_KEY)

    pedir    = int(cantidad * 1.5)
    ejemplos = "\n".join(f'  "{e}"' for e in EJEMPLOS_REALES.get(categoria, []))
    patron   = RESTRICCIONES_PATRON.get(categoria, "")

    excluir_bloque = ""
    if excluir:
        muestra = sorted(excluir)[:40]
        excluir_bloque = (
            "\nEXCLUSIONES — NO generes ninguna de estas oraciones ni variantes muy similares:\n"
            + "\n".join(f'  "{e}"' for e in muestra)
        )

    prompt_s = (
        "Eres un especialista en linguistica computacional y lenguas amazonicas "
        "del Peru. Tu tarea es generar oraciones en espanol para construir un "
        "corpus de clasificacion de intenciones para la lengua shiwilu (jebero, "
        "ISO 639-3: jeb), hablada en el distrito de Jeberos, Loreto, Peru."
    )
    prompt_u = f"""Genera exactamente {pedir} oraciones en espanol para \
la categoria de intencion: {categoria} - {DESC_INT[categoria]}

ESTILO OBLIGATORIO — imita estos ejemplos reales del corpus:
{ejemplos}

RESTRICCIONES:
1. Entre {LONG_MIN} y {LONG_MAX} palabras por oracion (las reales tienen 1-2 palabras en promedio).
2. Registro oral, cotidiano y conciso — NO oraciones largas ni literarias.
3. Usa vocabulario de los dominios del shiwilu cuando sea natural:
{vocab_str}
4. NO uses: internet, celular, computadora, dinero, carro, supermercado.
5. Cada oracion debe ser inequivocamente de la categoria {categoria}.
6. No repitas oraciones similares entre si.
{patron}
{excluir_bloque}

Responde UNICAMENTE con un JSON array valido, sin texto adicional:
[{{"oracion": "texto aqui", "dominio": "D1_Naturaleza"}}, ...]"""

    resp  = client.messages.create(
        model=MODELO, max_tokens=8192,
        system=prompt_s,
        messages=[{"role": "user", "content": prompt_u}],
    )
    texto = resp.content[0].text.strip()
    texto = re.sub(r"^```(?:json)?\s*", "", texto, flags=re.DOTALL)
    texto = re.sub(r"\s*```$",           "", texto, flags=re.DOTALL)
    todas = json.loads(texto)

    por_longitud = [
        item for item in todas
        if LONG_MIN <= len(item["oracion"].split()) <= LONG_MAX
    ]

    limpias = [
        item for item in por_longitud
        if _norm(item["oracion"]) not in excluir
    ]

    descartadas_long = len(todas) - len(por_longitud)
    descartadas_dup  = len(por_longitud) - len(limpias)
    if descartadas_long:
        print(f"    [longitud] {descartadas_long} descartadas por longitud")
    if descartadas_dup:
        print(f"    [duplicado] {descartadas_dup} descartadas por duplicado con pares anotados")

    return limpias[:cantidad]


def crear_excel(filas: list):
    wb = Workbook()

    ws = wb.active
    ws.title = "Para traducir"
    ws.append(["#", "ID", "Espanol", "Intencion", "Shiwilu (completar)"])
    aplicar_header(ws, "1B5E20")

    for i, row in enumerate(filas, 1):
        intent = row["intencion"]
        ws.append([i, row["id"], row["esp"], intent, ""])
        aplicar_fila(ws, COLOR_INT.get(intent, "FFFFFF"))
        ws.cell(ws.max_row, 5).fill = PatternFill("solid", fgColor="FFFFFF")

    for col, w in [("A",5),("B",14),("C",52),("D",12),("E",48)]:
        ancho(ws, col, w)
    ws.freeze_panes = "A2"

    ws2 = wb.create_sheet("Instrucciones")
    from openpyxl.styles import Font as F
    for texto, negrita in [
        ("INSTRUCCIONES PARA EL TRADUCTOR", True), ("", False),
        ("Completa la columna 'Shiwilu (completar)' con la traduccion al shiwilu.", False),
        ("", False),
        ("Traduce lo mas fielmente posible al shiwilu hablado en Jeberos.", False),
        ("Si no hay equivalente exacto, usa la expresion mas cercana.", False),
        ("Si una palabra no existe en shiwilu, escribe [sin equiv.]", False),
    ]:
        ws2.append([texto])
        ws2.cell(ws2.max_row, 1).font = F(name="Arial", bold=negrita, size=11)
    ancho(ws2, "A", 80)

    ws3 = wb.create_sheet("Resumen")
    ws3.append(["Intencion", "Descripcion", "Oraciones generadas"])
    aplicar_header(ws3, "1B5E20")
    conteo: dict = {}
    for row in filas:
        conteo[row["intencion"]] = conteo.get(row["intencion"], 0) + 1
    for intent, n in conteo.items():
        ws3.append([intent, DESC_INT.get(intent, ""), n])
        aplicar_fila(ws3, COLOR_INT.get(intent, "FFFFFF"))
    for col, w in [("A",12),("B",35),("C",16)]:
        ancho(ws3, col, w)

    ruta = INTERMEDIOS / "2_oraciones_para_traducir.xlsx"
    wb.save(ruta)
    print(f"  [OK] {ruta.name}  ({len(filas)} oraciones)")


if __name__ == "__main__":
    print("[2] Generando oraciones con API de Claude...")

    deficit  = leer_deficit()
    if not deficit:
        print("  Sin deficit — todas las categorias tienen 100 pares.")
        sys.exit(0)

    excluir   = leer_corpus_existente()
    vocab_str = vocab_para_prompt()

    print(f"  Categorias con deficit: {deficit}")
    print(f"  Oraciones excluidas del corpus: {len(excluir)}")
    print()

    filas  = []
    logs   = []
    conteo = {}

    for categoria, cantidad in deficit.items():
        print(f"  [{categoria}] Generando {cantidad} oraciones...")
        oraciones = generar_oraciones(categoria, cantidad, vocab_str, excluir)

        if len(oraciones) < cantidad:
            print(f"    [aviso] Se obtuvieron {len(oraciones)} de {cantidad} solicitadas")

        logs.append({
            "categoria": categoria,
            "solicitadas": cantidad,
            "obtenidas":   len(oraciones),
            "timestamp":   datetime.now().isoformat(),
        })
        for item in oraciones:
            conteo[categoria] = conteo.get(categoria, 0) + 1
            filas.append({
                "id":        f"{categoria}_NEW_{conteo[categoria]:03d}",
                "esp":       item["oracion"],
                "intencion": categoria,
                "dominio":   item.get("dominio", "D8_Social"),
            })
        print(f"  [OK] {len(oraciones)} oraciones limpias para {categoria}")

    with open(LOGS / "2_oraciones_generadas_log.json", "w", encoding="utf-8") as f:
        json.dump(logs, f, ensure_ascii=False, indent=2)

    crear_excel(filas)
    print()
    print(f"  Total: {len(filas)} oraciones sin duplicados, listas para traduccion")
